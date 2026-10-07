#!/usr/bin/env python3
"""Build data/metrics/attendance.json from NYC DOE End-of-Year Attendance and
Chronic Absenteeism Data (school level, 2018-19 through 2024-25).

Raw inputs
  - School file (88 MB, too big for the repo):
      public-school-attendance-results-2019-2025.xlsx
    Looked for at ATTENDANCE_XLSX (env var), then in the session scratchpad
    path below, then in $TMPDIR/nyc-schools-atlas-sources/. If none exists it
    is downloaded from FILE_URL into $TMPDIR/nyc-schools-atlas-sources/.
  - Citywide file (small, kept in the repo):
      data/sources/nyc_doe_eoy_attendance_citywide_2018-19_to_2024-25.xlsx
    Used only for the citywide reference values written into "notes".

Values taken: sheet "All Students", Grade == "All Grades",
Category == "All Students". Percent columns are divided by 100 and stored as
decimals rounded to 4 places. "s" (suppressed) becomes null.

Run from repo root:  python3 scripts/metrics/build_attendance.py
"""

import json
import os
import re
import statistics
import tempfile
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "metrics" / "attendance.json"
CITYWIDE = ROOT / "data" / "sources" / "nyc_doe_eoy_attendance_citywide_2018-19_to_2024-25.xlsx"

FILE_NAME = "public-school-attendance-results-2019-2025.xlsx"
FILE_URL = ("https://saintrafileprod01.blob.core.windows.net/prd-intra/docs/default-source/"
            "large-files/public-school-attendance-results-2019-2025.xlsx")
SOURCE_URL = ("https://infohub.nyced.org/reports/students-and-schools/school-quality/"
              "information-and-data-overview/end-of-year-attendance-and-chronic-absenteeism-data")
CITYWIDE_URL = ("https://infohub.nyced.org/docs/default-source/default-document-library/"
                "public-citywide-attendance-results-2019-2025.xlsx")
SCRATCH = Path("/private/tmp/claude-501/-Users-joshgreenman-Experiments/"
               "92d82c73-da9a-4cd9-bf8c-aaf5be19c7ce/scratchpad/attgrad") / FILE_NAME
CACHE = Path(tempfile.gettempdir()) / "nyc-schools-atlas-sources" / FILE_NAME

YEAR = "2024-25"
PREV_YEAR = "2023-24"
FETCHED = "2026-10-07"
SHEET = "All Students"
EXPECTED_HEADER = (
    "DBN", "School Name", "Grade", "Year", "Category", "# Total Days", "# Days Absent",
    "# Days Present", "% Attendance", "# Contributing 10+ Total Days and 1+ Pres Day",
    "# Chronically Absent", "% Chronically Absent",
)
MIN_SCHOOLS = 1500          # 2024-25 file has 1,563 "All Grades" school rows
SUPPRESSED = {"s"}
DBN_RE = re.compile(r"^\d{2}[MXKQR]\d{3}$")


def raw_path():
    env = os.environ.get("ATTENDANCE_XLSX")
    for p in [Path(env) if env else None, SCRATCH, CACHE]:
        if p and p.exists():
            return p
    CACHE.parent.mkdir(parents=True, exist_ok=True)
    print(f"Downloading {FILE_URL} -> {CACHE}")
    urllib.request.urlretrieve(FILE_URL, CACHE)
    return CACHE


def pct(v, marker_log):
    """Source percent (0-100) -> decimal 0-1; suppressed -> None."""
    if isinstance(v, str):
        if v.strip() in SUPPRESSED:
            marker_log.add(v.strip())
            return None
        raise ValueError(f"Unexpected text value in numeric cell: {v!r}")
    if v is None:
        return None
    return round(float(v) / 100.0, 4)


def count(v, marker_log):
    if isinstance(v, str):
        if v.strip() in SUPPRESSED:
            marker_log.add(v.strip())
            return None
        raise ValueError(f"Unexpected text value in count cell: {v!r}")
    if v is None:
        return None
    if float(v) != int(v):
        raise ValueError(f"Non-integer count: {v!r}")
    return int(v)


def citywide_reference():
    wb = openpyxl.load_workbook(CITYWIDE, read_only=True)
    ws = wb["All Students"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = rows[0]
    gi, ci, yi = hdr.index("Grade"), hdr.index("Category"), hdr.index("Year")
    ai, chi = hdr.index("% Attendance"), hdr.index("% Chronically Absent")
    out = {}
    for r in rows[1:]:
        if r[gi] == "All Grades" and r[ci] == "All Students":
            out[r[yi]] = (r[ai], r[chi])
    for y in (YEAR, PREV_YEAR):
        if y not in out:
            raise RuntimeError(f"Citywide file lacks All Grades/All Students row for {y}")
    return out


def main():
    src = raw_path()
    wb = openpyxl.load_workbook(src, read_only=True)
    ws = wb[SHEET]
    it = ws.iter_rows(values_only=True)
    header = tuple(next(it))
    if header != EXPECTED_HEADER:
        raise RuntimeError(f"Header changed in {SHEET!r}: {header}")
    H = {h: i for i, h in enumerate(header)}

    by_year = {YEAR: {}, PREV_YEAR: {}}
    markers = set()
    for r in it:
        if r[H["Grade"]] != "All Grades" or r[H["Category"]] != "All Students":
            continue
        yr = r[H["Year"]]
        if yr not in by_year:
            continue
        dbn = str(r[H["DBN"]]).strip().upper()
        if not DBN_RE.match(dbn):
            raise RuntimeError(f"Bad DBN {dbn!r}")
        if dbn in by_year[yr]:
            raise RuntimeError(f"Duplicate All Grades row for {dbn} {yr}")
        by_year[yr][dbn] = r

    cur, prev = by_year[YEAR], by_year[PREV_YEAR]
    if len(cur) < MIN_SCHOOLS or len(prev) < MIN_SCHOOLS:
        raise RuntimeError(f"Too few school rows: {YEAR}={len(cur)}, {PREV_YEAR}={len(prev)}")

    schools = {}
    for dbn, r in sorted(cur.items()):
        rec = {
            "attendance_rate": pct(r[H["% Attendance"]], markers),
            "chronic_absent": pct(r[H["% Chronically Absent"]], markers),
            "students_counted": count(r[H["# Contributing 10+ Total Days and 1+ Pres Day"]], markers),
            "chronic_absent_prev": None,
        }
        if dbn in prev:
            rec["chronic_absent_prev"] = pct(prev[dbn][H["% Chronically Absent"]], markers)
        schools[dbn] = rec

    # Sanity guards against an inverted or mis-mapped column.
    ca = [s["chronic_absent"] for s in schools.values() if s["chronic_absent"] is not None]
    ar = [s["attendance_rate"] for s in schools.values() if s["attendance_rate"] is not None]
    if len(ca) < MIN_SCHOOLS * 0.9 or len(ar) < MIN_SCHOOLS * 0.9:
        raise RuntimeError(f"Too many nulls: chronic={len(ca)}, attendance={len(ar)}")
    med = statistics.median(ca)
    if not (0.15 < med < 0.60):
        raise RuntimeError(f"Median school chronic absence {med:.3f} is implausible; check column mapping")
    pairs = [(s["attendance_rate"], s["chronic_absent"]) for s in schools.values()
             if s["attendance_rate"] is not None and s["chronic_absent"] is not None]
    hi_att = [c for a, c in pairs if a >= 0.95]
    lo_att = [c for a, c in pairs if a < 0.85]
    if hi_att and lo_att and statistics.median(hi_att) >= statistics.median(lo_att):
        raise RuntimeError("Chronic absence does not fall as attendance rises; column likely inverted")

    cw = citywide_reference()
    only_prev = sorted(set(prev) - set(cur))

    out = {
        "key": "attendance",
        "title": "Attendance and chronic absenteeism",
        "source": "NYC Department of Education (NYC Public Schools), End-of-Year Attendance and Chronic Absenteeism Data, school level",
        "source_url": SOURCE_URL,
        "file_url": FILE_URL,
        "vintage": YEAR,
        "fetched": FETCHED,
        "fields": {
            "attendance_rate": {
                "label": "Attendance rate",
                "definition": "Total days present divided by total days on register for all students at the school over the school year (all grades, including Pre-K where offered).",
                "source_field": "'% Attendance' (sheet 'All Students', Grade = 'All Grades', Category = 'All Students', Year = '2024-25')",
                "unit": "pct",
                "direction": "higher",
            },
            "chronic_absent": {
                "label": "Chronically absent",
                "definition": "Share of students with attendance of 90 percent or less (absent 10 percent or more of their days). Students must be enrolled at least 10 days and present at least 1 day to be counted.",
                "source_field": "'% Chronically Absent' (sheet 'All Students', Grade = 'All Grades', Category = 'All Students', Year = '2024-25')",
                "unit": "pct",
                "direction": "lower",
            },
            "chronic_absent_prev": {
                "label": "Chronically absent, 2023-24",
                "definition": "Same measure for the prior school year, 2023-24, from the same file.",
                "source_field": "'% Chronically Absent' (sheet 'All Students', Grade = 'All Grades', Category = 'All Students', Year = '2023-24')",
                "unit": "pct",
                "direction": "lower",
            },
            "students_counted": {
                "label": "Students counted",
                "definition": "Number of students enrolled at least 10 days and present at least 1 day, the denominator for the chronic absence rate.",
                "source_field": "'# Contributing 10+ Total Days and 1+ Pres Day' (sheet 'All Students', Grade = 'All Grades', Category = 'All Students', Year = '2024-25')",
                "unit": "count",
                "direction": "neutral",
            },
        },
        "notes": [
            f"Citywide reference (citywide file, All Grades, All Students): chronic absence {cw[YEAR][1]:.1f}% and attendance {cw[YEAR][0]:.1f}% in {YEAR}; chronic absence {cw[PREV_YEAR][1]:.1f}% and attendance {cw[PREV_YEAR][0]:.1f}% in {PREV_YEAR}. Citywide counts each student once, so it will not equal an average of school rates.",
            f"Median school-level chronic absence across the {len(ca)} schools with a value: {med * 100:.1f}%.",
            "Suppression marker: 's' (FERPA: rows with five or fewer students or 900 or fewer total days, plus complementary suppression; chronic absence also suppressed when five or fewer students contribute at least 20 days). Stored as null.",
            "Coverage per the DOE file notes: Districts 1-32 and 75. Charter schools, District 79 programs, home schooling and home/hospital instruction are excluded. Transfer schools are removed from the school-level file. Charter DBNs (84xxxx) will therefore have no attendance values.",
            "chronic_absent_prev is null where the school has no 2023-24 row or the 2023-24 value was suppressed.",
            f"{len(only_prev)} DBNs have a 2023-24 row but no 2024-25 row (closed, merged or renumbered); they are omitted.",
            "Attendance is attributed to the school the student attended at the time; a student enrolled at two schools in a year counts at both.",
            "Percent columns in the source are 0-100 with full float precision; stored here divided by 100 and rounded to 4 decimals.",
            f"Raw school file is 88 MB and is not committed. It is downloaded from file_url; the build script caches it at {CACHE}.",
        ],
        "schools": schools,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, indent=1, ensure_ascii=False) + "\n")
    n = lambda f: sum(1 for s in schools.values() if s[f] is not None)
    print(f"Read {src}")
    print(f"Wrote {OUT}: {len(schools)} schools; "
          + ", ".join(f"{f}={n(f)}" for f in ("attendance_rate", "chronic_absent", "chronic_absent_prev", "students_counted"))
          + f"; markers seen: {sorted(markers)}; median chronic {med:.3f}")


if __name__ == "__main__":
    main()
