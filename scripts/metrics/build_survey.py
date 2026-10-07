#!/usr/bin/env python3
"""Build data/metrics/survey.json from the 2026 NYC School Survey public data files.

Raw files (unmodified downloads, linked as "Download the 2026 K-12 Family, ... Student,
... K-12 Teacher ... survey data for all New York City public schools (Excel)" on
https://infohub.nyced.org/reports/students-and-schools/school-quality/nyc-school-survey):

  data/sources/nyc_doe_school_survey_2026_public_data_file_teacher.xlsx
      <- https://infohub.nyced.org/docs/default-source/default-document-library/2026-public-data-file-teacher.xlsx
  data/sources/nyc_doe_school_survey_2026_public_data_file_student.xlsx
      <- https://infohub.nyced.org/docs/default-source/default-document-library/2026-public-data-file-student.xlsx
  data/sources/nyc_doe_school_survey_2026_public_data_file_guardian.xlsx
      <- https://infohub.nyced.org/docs/default-source/default-document-library/2026-public-data-file-guardian.xlsx

Sheets used
  "<Group> Pos & Neg %"  question-level percent favorable / unfavorable. Three header rows:
                         row 1 = question text (on the first of each pair of columns),
                         row 2 = response-category label (e.g. "Agree/Strongly agree"),
                         row 3 = variable code (e.g. t_q29_P). Data start on row 4.
  "new total tab" (teacher) / "Total" (student, guardian)
                         one header row: response count, response rate and measure
                         ("... Score") percent-favorable scores, 0-100.

Every item column is located by its variable code and then checked against the question
wording in row 1 and the category label in row 2, so a shifted column raises instead of
silently writing the wrong question.

Cell handling: numbers are kept (item shares are already decimals 0-1, rounded by DOE to
two places; measure scores 0-100 are divided by 100). "N/A" (DOE suppression: fewer than
five responses for the respondent group, or the item was not asked of that school) and
"#N/A" (an Excel lookup error that appears in two cells of the guardian Total sheet) and
blank cells become null. Any other text raises.
"""
import json
import os
import re
from collections import Counter

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
SRC_DIR = os.path.join(ROOT, "data", "sources")
OUT = os.path.join(ROOT, "data", "metrics", "survey.json")

FETCHED = "2026-10-07"
BASE = "https://infohub.nyced.org/docs/default-source/default-document-library/"
FILES = {
    "teacher": ("nyc_doe_school_survey_2026_public_data_file_teacher.xlsx",
                BASE + "2026-public-data-file-teacher.xlsx"),
    "student": ("nyc_doe_school_survey_2026_public_data_file_student.xlsx",
                BASE + "2026-public-data-file-student.xlsx"),
    "guardian": ("nyc_doe_school_survey_2026_public_data_file_guardian.xlsx",
                 BASE + "2026-public-data-file-guardian.xlsx"),
}
# Row counts in the 2026 files: teacher Pos&Neg 1,761 / total 1,770; student 1,122 / 1,123;
# guardian 1,846 / 1,846. Fail well below those.
EXPECTED_MIN_ROWS = {"teacher": 1700, "student": 1050, "guardian": 1780}
DBN_RE = re.compile(r"^\d{2}[MXKQR]\d{3}$")
SUPPRESSION_MARKERS = {"N/A", "#N/A"}

# ---------------------------------------------------------------------------
# Question-level fields: (field, group, sheet, code, expected label, wording substring,
#                         label, definition, direction)
# ---------------------------------------------------------------------------
ITEMS = [
    ("teacher_principal_effective", "teacher", "Teacher Pos & Neg %", "t_q29_P",
     "Agree/Strongly agree",
     "q29. The principal/program leader at this school is an effective manager who makes the school run smoothly.",
     "Teachers: principal is an effective manager",
     "Share of responding teachers who agree or strongly agree: \"The principal/program "
     "leader at this school is an effective manager who makes the school run smoothly.\" (Teacher "
     "survey q29; part of the Teacher-Principal Trust measure.)",
     "higher"),
    ("teacher_recommend", "teacher", "Teacher Pos & Neg %", "t_q22_P",
     "Agree/Strongly agree",
     "q22. I would recommend this school to parents/guardians seeking a place for their child.",
     "Teachers who would recommend the school",
     "Share of responding teachers who agree or strongly agree: \"I would recommend this "
     "school to parents/guardians seeking a place for their child.\" (Teacher survey q22; School "
     "Commitment measure.)",
     "higher"),
    ("student_safe", "student", "Student Pos & Neg %", "s_q40_P",
     "Safe/Very Safe",
     "q40. Do you feel safe in the hallways and the cafeteria?",
     "Students who feel safe in hallways and cafeteria",
     "Share of responding students (grades 6-12) answering Safe or Very Safe to \"Do you feel safe "
     "in the hallways and the cafeteria?\" (Student survey q40; Safety measure.)",
     "higher"),
    ("student_supported", "student", "Student Pos & Neg %", "s_q13_P",
     "Most of your teachers/All of your teachers",
     "q13. How many of your teachers do you feel support you when you are upset?",
     "Students who feel supported by teachers",
     "Share of responding students answering Most or All of their teachers to \"How many of your "
     "teachers do you feel support you when you are upset?\" (Student survey q13; Student-Teacher "
     "Trust measure.)",
     "higher"),
    ("student_trusted_adult", "student", "Student Pos & Neg %", "s_q34_P",
     "Often/Very often",
     "q34. How often do you feel like there is at least one trusted adult at this school you could talk to if you needed to?",
     "Students with a trusted adult at school",
     "Share of responding students answering Often or Very often to \"How often do you feel like "
     "there is at least one trusted adult at this school you could talk to if you needed to?\" "
     "(Student survey q34; Social-Emotional measure.)",
     "higher"),
    ("student_vaping_often", "student", "Student Pos & Neg %", "s_q59_N",
     "Very often/Often",
     "q59. How often do students in this school vape?",
     "Students who say peers vape often",
     "Share of responding students answering Very often or Often to \"How often do students in "
     "this school vape?\" This is the unfavorable side of the item; DOE's percent favorable is the "
     "remainder (Sometimes/Never). New in 2026, reported under Additional Questions and not "
     "scored in any measure.",
     "lower"),
    ("parent_satisfied", "guardian", "Family Pos & Neg %", "p_q36_P",
     "Satisfied/Very satisfied",
     "q36. How satisfied are you with the following? The education my child has received this year.",
     "Families satisfied with the education",
     "Share of responding K-12 families answering Satisfied or Very satisfied to \"How satisfied "
     "are you with the following? The education my child has received this year.\" (Family "
     "survey q36; Family Satisfaction with Child's Education measure.)",
     "higher"),
]

# Total-sheet fields: (field, group, sheet, header, kind, label, definition, unit, direction)
TOTALS = [
    ("teacher_responses", "teacher", "new total tab", "Total Teacher Response Counts", "count",
     "Teacher responses",
     "Number of completed teacher surveys. Full-time teachers and staff employed at the school "
     "on Nov. 1, 2025 were eligible; DOE counts guidance counselors in the teacher population.",
     "count", "neutral"),
    ("teacher_response_rate", "teacher", "new total tab", "Total Teacher Response Rate", "rate",
     "Teacher response rate",
     "Completed teacher surveys divided by the eligible teacher population. For "
     "charter schools DOE estimates the population from enrollment or staffing.",
     "pct", "neutral"),
    ("student_responses", "student", "Total", "Total Student Response Count", "count",
     "Student responses",
     "Number of completed student surveys. Students in grades 6-12 enrolled by Oct. 1, 2025 "
     "were eligible; schools without those grades have no student survey.",
     "count", "neutral"),
    ("student_response_rate", "student", "Total", "Total Student Response Rate", "rate",
     "Student response rate",
     "Completed student surveys divided by students in grades 6-12 enrolled by Oct. 1, 2025.",
     "pct", "neutral"),
    ("parent_responses", "guardian", "Total", "Total Family Response Count", "count",
     "Family responses",
     "Number of completed K-12 family surveys. Families receive one survey per enrolled child.",
     "count", "neutral"),
    ("parent_response_rate", "guardian", "Total", "Total Family Response Rate", "rate",
     "Family response rate",
     "Completed family surveys divided by the number of enrolled students (one survey per child, "
     "students enrolled by Oct. 1, 2025).",
     "pct", "neutral"),
    # Measure scores (percent favorable, 0-100 in the file; stored 0-1 here)
    ("score_teacher_principal_trust", "teacher", "new total tab", "Teacher-Principal Trust Score",
     "score", "Teacher-principal trust (teacher survey)",
     "DOE measure score: the average percent favorable across the teacher items q28-q34 "
     "(respected by, trust, confidence in, effective manager, looks out for staff, puts "
     "children first, principal and APs work as a unit).",
     "pct", "higher"),
    ("score_instructional_leadership", "teacher", "new total tab", "Instructional Leadership Score",
     "score", "Instructional leadership (teacher survey)",
     "DOE measure score: the average percent favorable across teacher items q63-q70 on the "
     "principal and assistant principals setting expectations, tracking progress, giving "
     "feedback and planning instruction with teachers.",
     "pct", "higher"),
    ("score_student_safety", "student", "Total", "Safety Score", "score",
     "Safety (student survey)",
     "DOE measure score: the average percent favorable across student items q39, q40, q41 "
     "(feel safe in classes; hallways and cafeteria; locker rooms and bathrooms), q44 (School "
     "Safety Agents), q58 (physical fights), q60 (alcohol/drugs) and q61 (gang activity).",
     "pct", "higher"),
    ("score_student_teacher_trust", "student", "Total", "Student-Teacher Trust Score", "score",
     "Student-teacher trust (student survey)",
     "DOE measure score: the average percent favorable across student items q11, q12, q13, q27 "
     "and q28 (how many teachers are open to students' ideas, treat you with respect, support "
     "you when upset, treat students of different backgrounds and genders equally).",
     "pct", "higher"),
    ("score_family_satisfaction", "guardian", "Total", "Family Satisfaction with Child's Education Score",
     "score", "Family satisfaction with education (family survey)",
     "DOE measure score: the average percent favorable across family items q31, q32, q33, q36 "
     "and q37 (progress in reading and math, understanding grade-level work, satisfaction with "
     "the education and with the teachers).",
     "pct", "higher"),
    ("score_parent_principal_trust", "guardian", "Total", "Parent-Principal Trust Score", "score",
     "Parent-principal trust (family survey)",
     "DOE measure score: the average percent favorable across family items q5, q10 and q11 "
     "(respected by the principal, trust the principal at their word, principal is an "
     "effective manager).",
     "pct", "higher"),
]

FIELD_ORDER = [f[0] for f in ITEMS] + [f[0] for f in TOTALS]

suppressed = Counter()


def norm(s):
    return re.sub(r"\s+", " ", str(s).replace("\xa0", " ")).strip()


def cell(v, where):
    """Number, or None for a suppression marker / blank. Raises on any other text."""
    if v is None:
        return None
    if isinstance(v, bool):
        raise ValueError(f"unexpected boolean at {where}: {v!r}")
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if s == "":
        return None
    if s in SUPPRESSION_MARKERS:
        suppressed[s] += 1
        return None
    raise ValueError(f"unexpected text at {where}: {v!r}")


def dbn_of(v, where):
    d = str(v).strip().upper() if v is not None else ""
    if not DBN_RE.match(d):
        raise ValueError(f"bad DBN at {where}: {v!r}")
    return d


def load_sheet(group, sheet):
    path = os.path.join(SRC_DIR, FILES[group][0])
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"{FILES[group][0]}: sheet {sheet!r} missing; have {wb.sheetnames}")
    rows = [r for r in wb[sheet].iter_rows(values_only=True)]
    wb.close()
    return rows


def main():
    schools = {}

    def put(dbn, field, value):
        schools.setdefault(dbn, {f: None for f in FIELD_ORDER})[field] = value

    # ---------------- question-level items -----------------------------------------
    cache = {}
    for field, group, sheet, code, want_label, wording, *_ in ITEMS:
        key = (group, sheet)
        if key not in cache:
            cache[key] = load_sheet(group, sheet)
        rows = cache[key]
        qrow, lrow, crow = rows[0], rows[1], rows[2]
        if norm(lrow[0]) != "DBN":
            raise SystemExit(f"{sheet}: expected DBN in A2, got {lrow[0]!r}")
        codes = [norm(c) if c is not None else "" for c in crow]
        if codes.count(code) != 1:
            raise SystemExit(f"{sheet}: code {code} found {codes.count(code)} times")
        j = codes.index(code)
        if norm(lrow[j]) != want_label:
            raise SystemExit(f"{sheet} {code}: label {lrow[j]!r} != {want_label!r}")
        # question text sits on the first column of the N/P pair
        qtext = norm(qrow[j]) if qrow[j] else norm(qrow[j - 1])
        if norm(wording) not in qtext:
            raise SystemExit(f"{sheet} {code}: question text mismatch: {qtext!r}")
        data = [r for r in rows[3:] if r and r[0] is not None]
        if len(data) < EXPECTED_MIN_ROWS[group]:
            raise SystemExit(f"{sheet}: only {len(data)} school rows (< {EXPECTED_MIN_ROWS[group]})")
        for i, r in enumerate(data, start=4):
            dbn = dbn_of(r[0], f"{sheet}!A{i}")
            v = cell(r[j], f"{sheet} row {i} {code}")
            if v is not None:
                if not (0 <= v <= 1):
                    raise ValueError(f"{sheet} row {i} {code}: share out of range {v}")
                v = round(float(v), 4)
            put(dbn, field, v)

    # ---------------- total sheets -----------------------------------------------
    for field, group, sheet, header, kind, *_ in TOTALS:
        key = (group, sheet)
        if key not in cache:
            cache[key] = load_sheet(group, sheet)
        rows = cache[key]
        hdr = [norm(h) if h is not None else "" for h in rows[0]]
        if hdr[0] != "DBN":
            raise SystemExit(f"{sheet}: expected DBN in A1, got {hdr[0]!r}")
        if hdr.count(header) != 1:
            raise SystemExit(f"{FILES[group][0]} {sheet}: header {header!r} found {hdr.count(header)} times")
        j = hdr.index(header)
        data = [r for r in rows[1:] if r and r[0] is not None]
        if len(data) < EXPECTED_MIN_ROWS[group]:
            raise SystemExit(f"{sheet}: only {len(data)} school rows (< {EXPECTED_MIN_ROWS[group]})")
        for i, r in enumerate(data, start=2):
            dbn = dbn_of(r[0], f"{sheet}!A{i}")
            v = cell(r[j], f"{sheet} row {i} {header}")
            if v is not None:
                if kind == "count":
                    if float(v) != int(v) or v < 0:
                        raise ValueError(f"{sheet} row {i}: bad count {v}")
                    v = int(v)
                elif kind == "rate":
                    if v < 0:
                        raise ValueError(f"{sheet} row {i}: negative rate {v}")
                    v = round(float(v), 4)
                elif kind == "score":
                    if not (0 <= v <= 100):
                        raise ValueError(f"{sheet} row {i}: score out of range {v}")
                    v = round(float(v) / 100, 4)
            put(dbn, field, v)

    # ---------------- sanity checks ---------------------------------------------
    cov = {f: sum(1 for s in schools.values() if s[f] is not None) for f in FIELD_ORDER}
    min_cov = {
        "teacher_principal_effective": 1600, "teacher_recommend": 1600,
        "teacher_response_rate": 1700, "parent_response_rate": 1700,
        "parent_satisfied": 1500, "student_safe": 900, "student_response_rate": 1000,
    }
    for f, n in min_cov.items():
        if cov[f] < n:
            raise SystemExit(f"coverage too low for {f}: {cov[f]} < {n}")

    # Checks against figures THE CITY reported from the spring 2026 results (Oct. 7, 2026).
    s307 = schools.get("13K307", {})
    if s307.get("teacher_principal_effective") != 0.85:
        raise SystemExit(f"13K307 principal item {s307.get('teacher_principal_effective')} != 0.85")
    rec = [s["teacher_recommend"] for s in schools.values() if s["teacher_recommend"] is not None]
    n95 = sum(1 for v in rec if v >= 0.95)
    nhalf = sum(1 for v in rec if v <= 0.5)
    if not (520 <= n95 <= 560) or nhalf != 56:
        raise SystemExit(f"recommend checks off: {n95} schools >=95%, {nhalf} <=50%")

    fields = {}
    for field, group, sheet, code, want_label, wording, label, definition, direction in ITEMS:
        fields[field] = {
            "label": label,
            "definition": definition,
            "source_field": f"{code} (row-2 label \"{want_label}\") under \"{wording}\"; "
                            f"sheet \"{sheet}\", {FILES[group][0]}",
            "unit": "pct",
            "direction": direction,
            "respondents": group if group != "guardian" else "family",
            "file_url": FILES[group][1],
        }
    for field, group, sheet, header, kind, label, definition, unit, direction in TOTALS:
        fields[field] = {
            "label": label,
            "definition": definition,
            "source_field": f"\"{header}\"; sheet \"{sheet}\", {FILES[group][0]}"
                            + ("; file value is 0-100, divided by 100 here" if kind == "score" else ""),
            "unit": unit,
            "direction": direction,
            "respondents": group if group != "guardian" else "family",
            "file_url": FILES[group][1],
        }

    over1 = sorted(d for d, s in schools.items()
                   for f in ("teacher_response_rate", "student_response_rate", "parent_response_rate")
                   if s[f] is not None and s[f] > 1)

    out = {
        "key": "survey",
        "title": "NYC School Survey, 2026",
        "source": "NYC Public Schools (DOE), NYC School Survey 2026 public data files "
                  "(K-12 Teacher, Student and K-12 Family)",
        "source_url": "https://infohub.nyced.org/reports/students-and-schools/school-quality/nyc-school-survey",
        "file_url": [FILES[g][1] for g in ("teacher", "student", "guardian")],
        "vintage": "spring 2026 (administered Feb. 9 - Apr. 30, 2026; school year 2025-26)",
        "fetched": FETCHED,
        "fields": fields,
        "notes": [
            "Three separate DOE workbooks feed this file (teacher, student, K-12 family); file_url "
            "is therefore a list, and each field carries its own file_url.",
            "Percent favorable is DOE's own calculation: favorable responses divided by all "
            "favorable plus unfavorable responses to that question. 'I don't know', 'N/A' and "
            "blank answers are left out of the denominator (2026 Scoring Technical Guide). DOE "
            "rounds the shares to two decimals in the file.",
            "Suppression: the files print 'N/A' where a respondent group had fewer than five "
            "responses, or where a question did not apply to the school (e.g. elementary-only "
            "items). Those cells are null. Suppressed-cell counts in the columns used here: "
            + ", ".join(f"{k!r} x{v}" for k, v in sorted(suppressed.items())) + ". (The guardian "
            "Total sheet also has two '#N/A' Excel-error cells, in its Teacher Response Rate "
            "column, which is not used here; the script would null them too.)",
            "A null can also mean the school is simply not in that workbook: only schools serving "
            "grades 6-12 have a student survey, and a few District 75 schools appear in a Total "
            "sheet (response counts and scores) but not in the question-level sheet.",
            "Only the all-respondents school value is used. The workbooks hold no citywide row; "
            "demographic breakdowns are on separate sheets and are not used.",
            "The 2026 student survey was substantially rewritten. DOE says student results are a "
            "new baseline and should not be compared with earlier years. Teacher and family "
            "surveys had only minor revisions.",
            "Who counts as a teacher: DOE includes guidance counselors in the teacher population. "
            "Support staff (paraprofessionals, parent coordinators, social workers, etc.) also take "
            "a version of the survey, but the workbook does not say whether they are in these "
            "counts. Its response counts sum to 65,980, close to the 66,566 teacher respondents "
            "DOE's 2026 citywide deck reports excluding support staff, which suggests teachers "
            "only. The 2026 wording of the principal item is 'principal/program leader'; earlier "
            "years (and THE CITY's 2023 quote) used 'principal/school leader'.",
            "student_vaping_often is stored as the unfavorable share (students saying peers vape "
            "very often or often), so lower is better. It is an Additional Question that DOE does "
            "not score in any measure.",
            "student_safe uses q40 (hallways and cafeteria). The same survey also asks about "
            "feeling safe in classes (q39; 92% favorable citywide per DOE's 2026 deck) and in "
            "locker rooms and bathrooms (q41; 84%). DOE's deck puts q40 at 87% citywide.",
            "score_* fields are DOE measure scores from the workbooks' Total sheets: the average of "
            "the percent-favorable values of the measure's questions (checked here against the "
            "question-level sheets for sample schools). They are survey-only measures, not the "
            "School Quality Report framework ratings (Rigorous Instruction, Effective School "
            "Leadership, Trust, Supportive Environment, etc.), which are not in these files.",
            "Response rates come from each group's own workbook, as published (DOE rounds them to "
            "two decimals). For charter schools DOE estimates the teacher population. "
            f"Schools with any response rate above 100%: {len(set(over1))}.",
            "Checks against THE CITY's Oct. 7, 2026 guide: P.S. 307 (13K307) principal item = 85% "
            f"(matches); {n95} schools with 95%+ of teachers recommending (THE CITY: about 540); "
            f"{nhalf} schools where half or more of teachers did not agree they would recommend "
            "(THE CITY: 56). THE CITY's 87% citywide figure for the principal item could not be "
            "reproduced from this file: summing the file's own response counts gives about 86% "
            "(55,823 favorable of 65,033 favorable+unfavorable). The difference may be in how "
            "Panorama's citywide report aggregates; it is not a school-level value.",
            "2026 results are also on Panorama Education (secure.panoramaed.com/nycdoe), one page "
            "per school; this file uses only DOE's bulk Excel releases.",
        ],
        "schools": dict(sorted(schools.items())),
    }
    with open(OUT, "w") as f:
        json.dump(out, f, ensure_ascii=False, indent=1)
    print(f"wrote {OUT}: {len(schools)} schools")
    for f_ in FIELD_ORDER:
        print(f"  {f_:34s} {cov[f_]}")
    print("suppression markers:", dict(suppressed))


if __name__ == "__main__":
    main()
