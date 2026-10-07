#!/usr/bin/env python3
"""Build data/metrics/graduation.json.

Raw inputs
  1. NYC DOE Graduation Results, school level, Cohorts 2012-2021 (Classes of
     2016-2025): 2025-graduation-rates-public-school.xlsx (56 MB, not in repo).
     Looked for at GRAD_XLSX (env var), then the session scratchpad path below,
     then $TMPDIR/nyc-schools-atlas-sources/; downloaded from GRAD_URL if absent.
  2. NYC DOE 2024-25 School Quality Report, High School results:
     data/sources/nyc_doe_sqr_hs_2024-25.xlsx (byte-identical to
     https://infohub.nyced.org/docs/default-source/default-document-library/202425-hs-sqr-results.xlsx).
     Used only for the 4-year College and Career Readiness score.

Graduation values: sheet "All", Category "All Students".
  - 4-year fields: latest Cohort Year with "4 year August" rows (Cohort 2021).
  - 6-year fields: latest Cohort Year with "6 year August" rows (Cohort 2019).
  - Transfer-school rate: sheet "Transfer Schools", latest Cohort Year.
Percent columns are 0-100 in the source; stored as decimals rounded to 4 places.
"s" -> null.

Run from repo root:  python3 scripts/metrics/build_graduation.py
"""

import json
import os
import re
import statistics
import tempfile
import urllib.request
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "data" / "metrics" / "graduation.json"
SQR_HS = ROOT / "data" / "sources" / "nyc_doe_sqr_hs_2024-25.xlsx"
CITYWIDE = ROOT / "data" / "sources" / "nyc_doe_graduation_citywide_cohorts_2012-2021_class_of_2025.xlsx"

GRAD_NAME = "2025-graduation-rates-public-school.xlsx"
GRAD_URL = "https://infohub.nyced.org/docs/default-source/default-document-library/" + GRAD_NAME
GRAD_PAGE = "https://infohub.nyced.org/reports/academics/graduation-results"
SQR_URL = "https://infohub.nyced.org/docs/default-source/default-document-library/202425-hs-sqr-results.xlsx"
SQR_PAGE = ("https://infohub.nyced.org/reports/students-and-schools/school-quality/"
            "school-quality-reports-and-resources/school-quality-reports-citywide-results")
SCRATCH = Path("/private/tmp/claude-501/-Users-joshgreenman-Experiments/"
               "92d82c73-da9a-4cd9-bf8c-aaf5be19c7ce/scratchpad/attgrad") / GRAD_NAME
CACHE = Path(tempfile.gettempdir()) / "nyc-schools-atlas-sources" / GRAD_NAME
FETCHED = "2026-10-07"

GRAD_HEADER = (
    "DBN", "School Name", "Category", "Cohort Year", "Cohort", "# Total Cohort", "# Grads",
    "% Grads", "# Total Regents", "% Total Regents of Cohort", "% Total Regents of Grads",
    "# Advanced Regents", "% Advanced Regents of Cohort", "% Advanced Regents of Grads",
    "# Regents without Advanced", "% Regents without Advanced of Cohort",
    "% Regents without Advanced of Grads", "# Local", "% Local of Cohort", "% Local of Grads",
    "# Still Enrolled", "% Still Enrolled", "# Dropout", "% Dropout",
)
TRANSFER_HEADER = ("DBN", "School Name", "Category", "Cohort Year", "# Total Cohort", "# Grads", "% Grads")
SQR_DBN = "DBN"
SQR_CCR = "Metric Value - 4-Year College and Career Readiness - All Students"
SQR_CCR_N = "N count - 4-Year College and Career Readiness - All Students"

MIN_4YR = 450        # Cohort 2021 4-year August: 472 school rows
MIN_6YR = 450        # Cohort 2019 6-year August: 468 school rows
MIN_TRANSFER = 50    # Cohort 2019 transfer tab: 57 rows
MIN_CCR = 380        # 2024-25 HS SQR: 415 numeric 4-year CCR values
SUPPRESSED = {"s"}
SQR_SUPPRESSED = {"N<15", "N<10", ""}
DBN_RE = re.compile(r"^\d{2}[MXKQR]\d{3}$")


def grad_path():
    env = os.environ.get("GRAD_XLSX")
    for p in [Path(env) if env else None, SCRATCH, CACHE]:
        if p and p.exists():
            return p
    CACHE.parent.mkdir(parents=True, exist_ok=True)
    print(f"Downloading {GRAD_URL} -> {CACHE}")
    urllib.request.urlretrieve(GRAD_URL, CACHE)
    return CACHE


def num(v, markers, allowed, scale=None, as_int=False):
    if isinstance(v, str) or v is None:
        s = "" if v is None else v.strip()
        if s in allowed:
            markers.add(s if s else "(blank)")
            return None
        raise ValueError(f"Unexpected text value: {v!r}")
    if as_int:
        if float(v) != int(v):
            raise ValueError(f"Non-integer count {v!r}")
        return int(v)
    x = float(v)
    if scale:
        x = x / scale
    return round(x, 4)


def dbn_of(v):
    d = str(v).strip().upper()
    if not DBN_RE.match(d):
        raise RuntimeError(f"Bad DBN {v!r}")
    return d


def citywide_rate(cohort_year, cohort):
    """Citywide '% Grads' for All Students from the citywide file (for notes)."""
    ws = openpyxl.load_workbook(CITYWIDE, read_only=True)["All"]
    it = ws.iter_rows(values_only=True)
    h = next(it)
    C = {k: i for i, k in enumerate(h)}
    for r in it:
        if r[C["Category"]] == "All Students" and r[C["Cohort Year"]] == cohort_year and r[C["Cohort"]] == cohort:
            return r[C["% Grads"]]
    raise RuntimeError(f"No citywide row for {cohort_year} {cohort}")


def main():
    src = grad_path()
    wb = openpyxl.load_workbook(src, read_only=True)

    # ---- "All" sheet --------------------------------------------------------
    it = wb["All"].iter_rows(values_only=True)
    hdr = tuple(next(it))
    if hdr != GRAD_HEADER:
        raise RuntimeError(f"Graduation 'All' header changed: {hdr}")
    H = {h: i for i, h in enumerate(hdr)}
    rows = [r for r in it if r[H["Category"]] == "All Students"]
    cy4 = max(r[H["Cohort Year"]] for r in rows if r[H["Cohort"]] == "4 year August")
    cy6 = max(r[H["Cohort Year"]] for r in rows if r[H["Cohort"]] == "6 year August")
    if cy4 < 2021 or cy6 < 2019:
        raise RuntimeError(f"Unexpected latest cohorts: 4yr={cy4}, 6yr={cy6}")

    def pick(cohort_year, cohort):
        out = {}
        for r in rows:
            if r[H["Cohort Year"]] == cohort_year and r[H["Cohort"]] == cohort:
                d = dbn_of(r[H["DBN"]])
                if d in out:
                    raise RuntimeError(f"Duplicate {d} {cohort_year} {cohort}")
                out[d] = r
        return out

    r4 = pick(cy4, "4 year August")
    r6 = pick(cy6, "6 year August")
    if len(r4) < MIN_4YR or len(r6) < MIN_6YR:
        raise RuntimeError(f"Too few rows: 4yr={len(r4)}, 6yr={len(r6)}")

    # ---- "Transfer Schools" sheet --------------------------------------------
    it = wb["Transfer Schools"].iter_rows(values_only=True)
    th = tuple(next(it))
    if th != TRANSFER_HEADER:
        raise RuntimeError(f"Transfer header changed: {th}")
    T = {h: i for i, h in enumerate(th)}
    trows = [r for r in it if r[T["Category"]] == "All Students" and r[T["DBN"]]]
    cyt = max(r[T["Cohort Year"]] for r in trows)
    rt = {}
    for r in trows:
        if r[T["Cohort Year"]] == cyt:
            d = dbn_of(r[T["DBN"]])
            if d in rt:
                raise RuntimeError(f"Duplicate transfer row {d}")
            rt[d] = r
    if len(rt) < MIN_TRANSFER:
        raise RuntimeError(f"Too few transfer rows: {len(rt)}")

    # ---- 2024-25 HS School Quality Report: 4-year CCR ----------------------
    sw = openpyxl.load_workbook(SQR_HS, read_only=True)["Instruction and Performance"]
    srows = list(sw.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(srows[:10]) if r and SQR_DBN in r and SQR_CCR in r)
    sh = srows[hi]
    S = {h: i for i, h in enumerate(sh) if h}
    rc = {}
    for r in srows[hi + 1:]:
        if not r[S[SQR_DBN]]:
            continue
        d = dbn_of(r[S[SQR_DBN]])
        if d in rc:
            raise RuntimeError(f"Duplicate SQR row {d}")
        rc[d] = r

    # ---- assemble -------------------------------------------------------
    markers, sqr_markers = set(), set()
    schools = {}

    def rec(d):
        return schools.setdefault(d, {
            "grad_4yr": None, "cohort_4yr": None, "regents_4yr": None, "adv_regents_4yr": None,
            "dropout_4yr": None, "grad_6yr": None, "cohort_6yr": None,
            "transfer_grad": None, "transfer_cohort": None, "ccr_4yr": None, "ccr_4yr_n": None,
        })

    for d, r in r4.items():
        s = rec(d)
        s["grad_4yr"] = num(r[H["% Grads"]], markers, SUPPRESSED, 100)
        s["cohort_4yr"] = num(r[H["# Total Cohort"]], markers, SUPPRESSED, as_int=True)
        s["regents_4yr"] = num(r[H["% Total Regents of Cohort"]], markers, SUPPRESSED, 100)
        s["adv_regents_4yr"] = num(r[H["% Advanced Regents of Cohort"]], markers, SUPPRESSED, 100)
        s["dropout_4yr"] = num(r[H["% Dropout"]], markers, SUPPRESSED, 100)
    for d, r in r6.items():
        s = rec(d)
        s["grad_6yr"] = num(r[H["% Grads"]], markers, SUPPRESSED, 100)
        s["cohort_6yr"] = num(r[H["# Total Cohort"]], markers, SUPPRESSED, as_int=True)
    for d, r in rt.items():
        s = rec(d)
        s["transfer_grad"] = num(r[T["% Grads"]], markers, SUPPRESSED, 100)
        s["transfer_cohort"] = num(r[T["# Total Cohort"]], markers, SUPPRESSED, as_int=True)
    n_ccr = 0
    for d, r in rc.items():
        v = num(r[S[SQR_CCR]], sqr_markers, SQR_SUPPRESSED)
        n = num(r[S[SQR_CCR_N]], sqr_markers, SQR_SUPPRESSED, as_int=True)
        if v is None and d not in schools:
            continue          # don't create empty records for SQR-only schools
        s = rec(d)
        s["ccr_4yr"] = v
        s["ccr_4yr_n"] = n
        if v is not None:
            n_ccr += 1
            if not (0 <= v <= 100):
                raise RuntimeError(f"CCR out of 0-100 range for {d}: {v}")
    if n_ccr < MIN_CCR:
        raise RuntimeError(f"Too few CCR values: {n_ccr}")

    g4 = [s["grad_4yr"] for s in schools.values() if s["grad_4yr"] is not None]
    if len(g4) < MIN_4YR * 0.9:
        raise RuntimeError(f"Too many null 4-year rates: {len(g4)}")
    if not (0.6 < statistics.median(g4) < 1.0):
        raise RuntimeError(f"Median 4-year grad rate {statistics.median(g4)} implausible")

    schools = dict(sorted(schools.items()))
    out = {
        "key": "graduation",
        "title": "Graduation, Regents diplomas and college and career readiness",
        "source": "NYC Department of Education (NYC Public Schools), Graduation Results, school level (Cohorts 2012-2021); 4-year CCR score from the 2024-25 School Quality Report, High School results",
        "source_url": GRAD_PAGE,
        "file_url": GRAD_URL,
        "vintage": f"Cohort {cy4} (Class of {cy4 + 4}), 4-year August; 6-year rate Cohort {cy6}; CCR from 2024-25 School Quality Report",
        "fetched": FETCHED,
        "fields": {
            "grad_4yr": {
                "label": f"4-year graduation rate (Class of {cy4 + 4})",
                "definition": f"Share of the Cohort {cy4} (students who first entered 9th grade in {cy4}-{str(cy4 + 1)[2:]}) who earned a Local or Regents diploma within four years, including August graduates.",
                "source_field": f"'% Grads' (sheet 'All', Category = 'All Students', Cohort Year = {cy4}, Cohort = '4 year August')",
                "unit": "pct", "direction": "higher",
            },
            "cohort_4yr": {
                "label": "4-year cohort size",
                "definition": f"Number of students in the school's Cohort {cy4} graduation cohort.",
                "source_field": f"'# Total Cohort' (sheet 'All', Category = 'All Students', Cohort Year = {cy4}, Cohort = '4 year August')",
                "unit": "count", "direction": "neutral",
            },
            "regents_4yr": {
                "label": "Regents diploma rate (4-year)",
                "definition": "Share of the 4-year cohort who earned a Regents diploma (with or without Advanced designation).",
                "source_field": f"'% Total Regents of Cohort' (sheet 'All', Category = 'All Students', Cohort Year = {cy4}, Cohort = '4 year August')",
                "unit": "pct", "direction": "higher",
            },
            "adv_regents_4yr": {
                "label": "Advanced Regents diploma rate (4-year)",
                "definition": "Share of the 4-year cohort who earned a Regents diploma with Advanced designation.",
                "source_field": f"'% Advanced Regents of Cohort' (sheet 'All', Category = 'All Students', Cohort Year = {cy4}, Cohort = '4 year August')",
                "unit": "pct", "direction": "higher",
            },
            "dropout_4yr": {
                "label": "Dropout rate (4-year)",
                "definition": "Share of the 4-year cohort recorded as dropouts as of August of the fourth year.",
                "source_field": f"'% Dropout' (sheet 'All', Category = 'All Students', Cohort Year = {cy4}, Cohort = '4 year August')",
                "unit": "pct", "direction": "lower",
            },
            "grad_6yr": {
                "label": f"6-year graduation rate (Cohort {cy6})",
                "definition": f"Share of the Cohort {cy6} (entered 9th grade in {cy6}-{str(cy6 + 1)[2:]}) who earned a Local or Regents diploma within six years, including August graduates. This is an older cohort than the 4-year rate.",
                "source_field": f"'% Grads' (sheet 'All', Category = 'All Students', Cohort Year = {cy6}, Cohort = '6 year August')",
                "unit": "pct", "direction": "higher",
            },
            "cohort_6yr": {
                "label": "6-year cohort size",
                "definition": f"Number of students in the school's Cohort {cy6} graduation cohort.",
                "source_field": f"'# Total Cohort' (sheet 'All', Category = 'All Students', Cohort Year = {cy6}, Cohort = '6 year August')",
                "unit": "count", "direction": "neutral",
            },
            "transfer_grad": {
                "label": "Transfer-school graduation rate",
                "definition": "DOE's alternative graduation rate for transfer high schools: share of the transfer school's graduation cohort (students whose transfer-school graduation deadline, end of year six or seven of high school, fell in 2025, plus earlier-deadline students who graduated in 2025) who earned a Regents or Local diploma. Not comparable to the 4-year rate.",
                "source_field": f"'% Grads' (sheet 'Transfer Schools', Category = 'All Students', Cohort Year = {cyt})",
                "unit": "pct", "direction": "higher",
            },
            "transfer_cohort": {
                "label": "Transfer-school cohort size",
                "definition": "Number of students in the transfer school's graduation cohort.",
                "source_field": f"'# Total Cohort' (sheet 'Transfer Schools', Category = 'All Students', Cohort Year = {cyt})",
                "unit": "count", "direction": "neutral",
            },
            "ccr_4yr": {
                "label": "College and career readiness score (4-year, 0-100)",
                "definition": "Average College and Career Readiness (CCR) score of students in the school's four-year cohort (Class of 2025) after their fourth year. Each student scores 0-100 based on course grades, test scores, advanced courses and endorsements. This is an average score, not the percent of students who are ready.",
                "source_field": f"'{SQR_CCR}' (2024-25 HS School Quality Report, sheet 'Instruction and Performance')",
                "unit": "number", "direction": "higher",
            },
            "ccr_4yr_n": {
                "label": "Students in CCR cohort",
                "definition": "Number of students contributing to the 4-year CCR score.",
                "source_field": f"'{SQR_CCR_N}' (2024-25 HS School Quality Report, sheet 'Instruction and Performance')",
                "unit": "count", "direction": "neutral",
            },
        },
        "notes": [
            f"Citywide reference (citywide graduation file, All Students): 4-year August rate {citywide_rate(cy4, '4 year August'):.1f}% for Cohort {cy4}; 6-year August rate {citywide_rate(cy6, '6 year August'):.1f}% for Cohort {cy6}. Citywide cohorts include out-of-district placements, so they do not sum from school rows.",
            f"Graduation file: '{GRAD_NAME}' from {GRAD_PAGE}. Latest 4-year cohort is Cohort {cy4} (Class of {cy4 + 4}); latest 6-year August cohort is Cohort {cy6} (Class of {cy6 + 4}, measured by August {cy6 + 6}).",
            "Graduation file coverage: school-level results are Districts 1-32 only. Charter schools are not on the 'All' tab, so charters have no grad_4yr/grad_6yr (InfoHub says a separate charter graduation report 'will be posted shortly'; it was not posted as of 2026-10-07). District 75 and District 79 schools have no school rows.",
            "Graduation suppression marker: 's' (rows with fewer than 5 students, plus complementary suppression). Stored as null.",
            f"Transfer schools: {len(rt)} DBNs on the 'Transfer Schools' tab for Cohort Year {cyt}; {sum(1 for d in rt if d in r4)} of them also have a regular 4-year row on the 'All' tab, and {sum(1 for d in rt if d.startswith('84'))} are charter transfer schools (the only charter rows in this file). DOE publishes the transfer rate because transfer schools enroll over-age, under-credited students; for those schools transfer_grad is the fairer measure.",
            "Transfer tab suppresses cohorts of 10 or fewer students (to align with the School Quality Report); its percentages are rounded in the source (e.g. 73, 73.3).",
            "regents_4yr includes Advanced Regents. Graduates = Regents + Local diplomas.",
            "CCR source: 2024-25 High School School Quality Report (data/sources/nyc_doe_sqr_hs_2024-25.xlsx = " + SQR_URL + ", landing page " + SQR_PAGE + "). The SQR attributes students to the last diploma-granting school as of June 30 of year four, so its cohorts differ slightly from the graduation file's.",
            "CCR suppression marker: 'N<15' (value) with blank N count. Stored as null. Charter high schools appear in the SQR file but have no CCR value.",
            "The 2024-25 Educator Guide notes CCR replaced the College Readiness Index (reported 2012-2020) and was introduced in the 2023-24 reports.",
            "Percent columns are 0-100 in the source; stored as decimals rounded to 4 places. ccr_4yr stays on its native 0-100 scale.",
            f"Raw graduation school file is 56 MB and is not committed. The build script downloads it from file_url and caches it at {CACHE}.",
        ],
        "schools": schools,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, indent=1, ensure_ascii=False) + "\n")
    fields = list(out["fields"])
    cov = {f: sum(1 for s in schools.values() if s[f] is not None) for f in fields}
    print(f"Read {src} and {SQR_HS}")
    print(f"Wrote {OUT}: {len(schools)} schools; coverage {cov}; markers {sorted(markers)} / SQR {sorted(sqr_markers)}")


if __name__ == "__main__":
    main()
